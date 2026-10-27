#!/usr/bin/env python3
"""
gen_testcases.py —— 从 docs/test-cases.json 生成 docs/test-cases.xlsx
表头固定：ID｜Category｜Scenario｜Precondition｜Request｜Expected Result｜Error Code
用法：python scripts/gen_testcases.py
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "test-cases.json"
OUT = ROOT / "docs" / "test-cases.xlsx"

data = json.loads(SRC.read_text(encoding="utf-8"))
header = data["header"]
rows = data["rows"]

wb = Workbook()
ws = wb.active
ws.title = data["sheet"]

ws.append(header)
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="4472C4")
    cell.alignment = Alignment(horizontal="center", vertical="center")

for row in rows:
    ws.append(row)

# 列宽与自动换行（内容列给足宽度，长文本不横向拉爆表格）
widths = [14, 10, 26, 30, 34, 46, 14]
for i, width in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = width
for row_cells in ws.iter_rows(min_row=2):
    for cell in row_cells:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

ws.freeze_panes = "A2"

wb.save(OUT)
print(f"已生成 {OUT}（{len(rows)} 条用例）")
